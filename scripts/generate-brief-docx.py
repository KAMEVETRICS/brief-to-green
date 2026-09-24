#!/usr/bin/env python3
"""Generate briefs/sample-feature-brief.docx for Bob document-understanding demo.

Prefers python-docx. If unavailable, decodes briefs/sample-feature-brief.docx.b64
(checked in for MCP/text-only push workflows).
"""

from __future__ import annotations

import base64
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "briefs" / "sample-feature-brief.docx"
B64 = ROOT / "briefs" / "sample-feature-brief.docx.b64"


def write_from_b64() -> None:
    parts_dir = ROOT / "briefs" / "docx-parts"
    if B64.exists():
        raw = B64.read_text().strip()
    elif parts_dir.exists():
        raw = "".join(p.read_text().strip() for p in sorted(parts_dir.glob("brief.part*.b64")))
    else:
        raise SystemExit(f"Missing {B64} (or briefs/docx-parts) and python-docx unavailable")
    OUT.write_bytes(base64.b64decode(raw))
    print(f"Decoded {OUT}")


def write_with_docx() -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading("Product brief: Priority filter & Due-date badges", 0)

    meta = doc.add_paragraph()
    meta.add_run("Product: ").bold = True
    meta.add_run("Sample Board (apps/sample-board)\n")
    meta.add_run("Author: ").bold = True
    meta.add_run("KAMEVETRICS (hackathon sample PRD)\n")
    meta.add_run("Status: ").bold = True
    meta.add_run("Ready for Bob Plan → Agent\n")
    meta.add_run("Priority: ").bold = True
    meta.add_run("P0 for Brief→Green demo")

    doc.add_heading("Context", 1)
    doc.add_paragraph(
        "Sample Board is a Kanban-style task board used as the real sample project "
        "for Brief→Green. Seed data already includes priority and dueDate. The "
        "intentional gap is UI: Priority filter and Due-date badges are not implemented."
    )

    doc.add_heading("Problem", 1)
    for item in (
        "Users cannot filter the board to High priority work.",
        "Overdue and soon deadlines are invisible on cards.",
        "Triage stays slower than standup needs.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Goals", 1)
    for item in (
        "Keep priority (low|medium|high) on every task — already seeded.",
        "Keep optional dueDate (YYYY-MM-DD) — already seeded.",
        "Add Priority filter above the board: All / High / Medium / Low.",
        "Show due-date badge on each card: upcoming · soon (≤2 days) · overdue.",
        "Cover with lightweight tests or a Bob-run QA checklist.",
    ):
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("Non-goals", 1)
    doc.add_paragraph(
        "Calendar picker redesign; recurring tasks; backend persistence; auth / multiplayer."
    )

    doc.add_heading("Acceptance criteria", 1)
    for item in (
        "Priority filter above the board; All shows everything; levels hide non-matches.",
        "Cards with dueDate show a badge with tone overdue / soon / upcoming.",
        "Cards without due date: no badge (or muted No due).",
        "Filter + badges compose correctly.",
        "npm run build green for sample-board.",
        "At least one automated test or Bob-run manual QA script for the happy path.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Technical hints for Bob Plan", 1)
    doc.add_paragraph(
        "Stack: Vite + React + TypeScript. Task already has priority/dueDate in "
        "src/types.ts and src/data.ts. Add PriorityFilter, DueBadge, helpers "
        "dueTone(date) and matchesPriority(task, filter). Files: src/App.tsx, "
        "src/components/*."
    )

    doc.add_heading("Success metric", 1)
    doc.add_paragraph(
        "Wall-clock from brief → Bob Plan to filter + badges + green checks, "
        "with session summaries saved under docs/bob-sessions/."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    # Keep text-safe companion for repos that cannot push binary via API.
    B64.write_text(base64.b64encode(OUT.read_bytes()).decode("ascii") + "\n")
    print(f"Wrote {OUT} and {B64}")


def main() -> None:
    try:
        write_with_docx()
    except ImportError:
        write_from_b64()


if __name__ == "__main__":
    main()
